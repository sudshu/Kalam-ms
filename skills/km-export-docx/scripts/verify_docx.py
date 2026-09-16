#!/usr/bin/env python3
"""Verify a rendered DOCX is well-formed and complete.

Checks: embedded images, unresolved citations, bibliography, raw LaTeX,
tables, file size, line numbers, double spacing, footer page numbers, and
figure-caption separators.
Prints results as JSON to stdout.
"""

import argparse
import json
import os

from docx import Document
from docx.shared import Inches
from docx.oxml.ns import qn


def verify_docx(
    filepath,
    expected_min_images=0,
    cover_letter=False,
    expect_full_width_images=False,
    expected_line_spacing=2.0,
):
    """Verify a rendered DOCX is well-formed and complete."""
    doc = Document(filepath)
    issues = []

    # 1. Count embedded images
    image_count = 0
    image_widths = []
    for para in doc.paragraphs:
        for run in para.runs:
            drawings = run._element.findall(qn('w:drawing'))
            image_count += len(drawings)
            for drawing in drawings:
                image_widths.extend(
                    int(ext.get('cx', 0))
                    for ext in drawing.findall('.//' + qn('wp:extent'))
                )
    if image_count < expected_min_images:
        issues.append(f"Expected at least {expected_min_images} images, found {image_count}")

    # 2. Check for unresolved pandoc citations [@...]
    raw_cites = []
    for para in doc.paragraphs:
        if '[@' in para.text:
            snippet = para.text[:80].replace('\n', ' ')
            raw_cites.append(snippet)
    if raw_cites:
        issues.append(f"Raw [@...] citations found in {len(raw_cites)} paragraphs: {raw_cites[0]}...")

    # 3. Check for a references/bibliography section
    has_bibliography = any(
        para.style.name == 'Bibliography' or
        (para.style.name.startswith('Heading') and 'references' in para.text.lower())
        for para in doc.paragraphs
    )
    bib_entries = sum(1 for p in doc.paragraphs if p.style.name == 'Bibliography')

    # 4. Check that LaTeX math didn't render as raw markup
    raw_latex = []
    for para in doc.paragraphs:
        text = para.text
        if '\\mathrm' in text or '\\frac' in text or '\\mathbf' in text:
            raw_latex.append(text[:60])
    if raw_latex:
        issues.append(f"Raw LaTeX found in {len(raw_latex)} paragraphs (math may not have rendered)")

    # 5. Check tables rendered
    table_count = len(doc.tables)

    # 6. File size sanity check (figures should add bulk)
    file_size_mb = os.path.getsize(filepath) / (1024 * 1024)

    # 7. Check line numbers are configured
    has_line_numbers = False
    for section in doc.sections:
        if section._sectPr.findall(qn('w:lnNumType')):
            has_line_numbers = True
            break

    # 8. Check figure-caption separators exist
    separator_count = 0
    for para in doc.paragraphs:
        pPr = para._element.find(qn('w:pPr'))
        if pPr is not None:
            pBdr = pPr.find(qn('w:pBdr'))
            if pBdr is not None and pBdr.find(qn('w:bottom')) is not None:
                separator_count += 1

    # 9. Body paragraphs must carry the spacing the export applied - double by default, or
    #    whatever --line-spacing was passed to format_docx.py. Tables remain compact.
    checked_spacing = [
        para for para in doc.paragraphs
        if para.text.strip()
    ]
    double_spaced_count = sum(
        1 for para in checked_spacing
        if para.paragraph_format.line_spacing == expected_line_spacing
    )
    double_spaced = bool(checked_spacing) and double_spaced_count == len(checked_spacing)

    # 10. Every section must expose a PAGE field in its footer.
    page_number_sections = 0
    for section in doc.sections:
        if any(
            (node.text or "").strip() == "PAGE"
            for node in section.footer._element.iter(qn("w:instrText"))
        ):
            page_number_sections += 1
    has_page_numbers = bool(doc.sections) and page_number_sections == len(doc.sections)

    if not cover_letter:
        if not double_spaced:
            issues.append(
                f"Expected line spacing {expected_line_spacing} in all text paragraphs; found "
                f"{double_spaced_count}/{len(checked_spacing)}"
            )
        if not has_page_numbers:
            issues.append(
                f"Expected footer page numbers in every section; found "
                f"{page_number_sections}/{len(doc.sections)}"
            )

    section = doc.sections[0]
    page_width = section.page_width or Inches(8.5)
    text_width = int(page_width - section.left_margin - section.right_margin)
    width_tolerance = max(10000, int(text_width * 0.005))
    full_width_images = sum(
        1 for width in image_widths
        if abs(width - text_width) <= width_tolerance
    )
    if expect_full_width_images and full_width_images != image_count:
        issues.append(
            f"Expected all {image_count} images at full text width; found "
            f"{full_width_images}"
        )

    result = {
        'file': os.path.basename(filepath),
        'size_mb': round(file_size_mb, 1),
        'paragraphs': len(doc.paragraphs),
        'images': image_count,
        'full_width_images': full_width_images,
        'text_width_cm': round(text_width / 360000, 2),
        'tables': table_count,
        'bibliography_entries': bib_entries,
        'has_bibliography': has_bibliography,
        'raw_citations': len(raw_cites),
        'raw_latex': len(raw_latex),
        'line_numbers': has_line_numbers,
        'expected_line_spacing': expected_line_spacing,
        'spacing_as_expected': double_spaced,
        'paragraphs_at_expected_spacing': double_spaced_count,
        'checked_spacing_paragraphs': len(checked_spacing),
        'page_numbers': has_page_numbers,
        'page_number_sections': page_number_sections,
        'sections': len(doc.sections),
        'figure_caption_separators': separator_count,
        'issues': issues,
    }
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(
        description='Verify a rendered DOCX is well-formed and complete.'
    )
    parser.add_argument('filepath', help='Path to the .docx file to verify')
    parser.add_argument(
        '--expected-images',
        type=int,
        default=0,
        metavar='N',
        help='Minimum number of expected embedded images (default: 0)'
    )
    parser.add_argument(
        '--cover-letter',
        action='store_true',
        help='Do not require manuscript double spacing or footer page numbers'
    )
    parser.add_argument(
        '--expect-full-width-images',
        action='store_true',
        help='Require every embedded image to span the usable text width'
    )
    parser.add_argument(
        '--line-spacing',
        type=float,
        default=2.0,
        help='Line spacing the export applied; must match format_docx.py --line-spacing (default 2.0)'
    )
    args = parser.parse_args()

    result = verify_docx(
        args.filepath,
        expected_min_images=args.expected_images,
        cover_letter=args.cover_letter,
        expect_full_width_images=args.expect_full_width_images,
        expected_line_spacing=args.line_spacing,
    )
    print(json.dumps(result, indent=2))
